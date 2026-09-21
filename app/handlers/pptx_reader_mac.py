"""Презентации на macOS: python-pptx. PowerPoint COM здесь нет."""
from __future__ import annotations

from pathlib import Path


def pptx_to_pdf(input_path, output_folder="convert_pdf"):
    """Достаёт текст из `.pptx`. Имя функции историческое (PDF на Mac не нужен)."""
    ext = Path(input_path).suffix.lower()
    if ext == ".ppt":
        raise RuntimeError(
            "Старый формат .ppt на macOS не читается. Сохраните файл как .pptx."
        )
    if ext != ".pptx":
        raise RuntimeError(f"Ожидался .ppt или .pptx, получен {ext or 'файл без расширения'}")

    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as e:
        raise RuntimeError("Установите python-pptx для чтения презентаций на macOS") from e

    parts: list[str] = []
    presentation = Presentation(input_path)
    for i, slide in enumerate(presentation.slides, 1):
        bits: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = (shape.text or "").strip()
                if text:
                    bits.append(text)
            elif shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        bits.append(" | ".join(cells))
        if bits:
            parts.append(f"Слайд {i}\n" + "\n".join(bits))
    text = "\n\n".join(parts)
    if not text.strip():
        raise RuntimeError("В презентации нет текстовых блоков")
    return text
