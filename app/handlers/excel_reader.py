from __future__ import annotations

import math
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime

import pandas as pd

LARGE_SHEET_ROWS = 100
EXCEL_EXTENSIONS = {".xlsx", ".xls", ".xlsm", ".xlsb", ".xltx", ".xltm"}
CSV_EXTENSIONS = {".csv"}
TABLE_EXTENSIONS = EXCEL_EXTENSIONS | CSV_EXTENSIONS
EXCEL_GRAPH_STUB = "[excel-graph]"
_SKIP_VALUES = {
    "",
    "-",
    "—",
    "n/a",
    "na",
    "нет данных",
    "не применимо",
    "none",
    "null",
}
_BOOLISH = {"да", "нет", "ложь", "истина", "true", "false", "0", "1"}
_SLASHY = re.compile(r"^[\\/]+$")


@dataclass
class ExcelTable:
    sheet_name: str
    headers: list[str]
    kinds: list[str]
    rows: list[list[str]]
    values_by_column: dict[str, list[str]] = field(default_factory=dict)


def is_excel_path(path: str) -> bool:
    return os.path.splitext(path or "")[1].lower() in EXCEL_EXTENSIONS


def is_csv_path(path: str) -> bool:
    return os.path.splitext(path or "")[1].lower() in CSV_EXTENSIONS


def is_table_path(path: str) -> bool:
    return os.path.splitext(path or "")[1].lower() in TABLE_EXTENSIONS


def split_excel(file_path: str) -> tuple[str, list[ExcelTable]]:
    """Короткие листы — markdown, длиннее 100 строк — табличная модель для графа."""
    if is_csv_path(file_path):
        return _split_csv(file_path)
    book = pd.ExcelFile(file_path)
    try:
        merges = _merged_ranges(file_path)
        small: list[str] = []
        tables: list[ExcelTable] = []
        for sheet_name in book.sheet_names:
            frame = book.parse(sheet_name, header=None)
            sheet_merges = merges.get(sheet_name) or []
            if _sheet_row_count(frame) > LARGE_SHEET_ROWS:
                table = _large_sheet(frame, sheet_merges, str(sheet_name))
                if table and table.rows:
                    tables.append(table)
                continue
            text = _sheet_text(frame, sheet_merges)
            if text:
                small.append(f"Лист: {sheet_name}\n{text}")
        return "\n\n".join(small), tables
    finally:
        book.close()


def excel_fingerprint(small_text: str, tables: list[ExcelTable]) -> str:
    parts = [small_text.strip()] if (small_text or "").strip() else []
    for table in tables:
        headers = "|".join(table.headers)
        parts.append(f"{EXCEL_GRAPH_STUB}\t{table.sheet_name}\t{len(table.rows)}\t{headers}")
    return "\n".join(parts)


def read_excel(file_path):
    """Текст коротких листов плюс метка длинных — чтобы digest и ingest видели файл."""
    small, tables = split_excel(file_path)
    return excel_fingerprint(small, tables)


def _split_csv(file_path: str) -> tuple[str, list[ExcelTable]]:
    frame = _read_csv_frame(file_path)
    sheet_name = os.path.splitext(os.path.basename(file_path))[0] or "csv"
    if _sheet_row_count(frame) > LARGE_SHEET_ROWS:
        table = _large_sheet(frame, [], sheet_name)
        if table and table.rows:
            return "", [table]
        return "", []
    text = _sheet_text(frame, [])
    if text:
        return f"Таблица: {sheet_name}\n{text}", []
    return "", []


def _read_csv_frame(file_path: str) -> pd.DataFrame:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            return pd.read_csv(
                file_path,
                header=None,
                encoding=encoding,
                sep=None,
                engine="python",
            )
        except Exception as error:
            last_error = error
    raise ValueError(f"Не удалось прочитать CSV: {last_error}") from last_error


def strip_excel_graph_stub(text: str) -> str:
    lines = []
    for line in (text or "").splitlines():
        if line.startswith(EXCEL_GRAPH_STUB):
            continue
        lines.append(line)
    return "\n".join(lines).strip()


def _merged_ranges(file_path) -> dict[str, list[tuple[int, int, int, int]]]:
    """Прямоугольники объединённых ячеек. Для .xls их нет: openpyxl их не читает."""
    if not str(file_path).lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        return {}
    try:
        import openpyxl
    except ImportError:
        return {}
    found: dict[str, list[tuple[int, int, int, int]]] = {}
    workbook = openpyxl.load_workbook(file_path, read_only=False, data_only=False)
    try:
        for name in workbook.sheetnames:
            sheet = workbook[name]
            found[name] = [
                (item.min_row - 1, item.min_col - 1, item.max_row - 1, item.max_col - 1)
                for item in sheet.merged_cells.ranges
            ]
    finally:
        workbook.close()
    return found


def _sheet_row_count(frame: pd.DataFrame) -> int:
    if frame is None or frame.empty:
        return 0
    return int(len(frame))


def _skip_value(text: str) -> bool:
    value = " ".join((text or "").split())
    if not value or _SLASHY.match(value):
        return True
    return value.casefold() in _SKIP_VALUES


def _skip_header(text: str) -> bool:
    value = " ".join((text or "").split())
    if not value or _SLASHY.match(value.replace(" ", "")):
        return True
    return False


def _unique_headers(raw: list[str]) -> list[str]:
    used: dict[str, int] = {}
    headers = []
    for index, item in enumerate(raw, 1):
        name = " ".join((item or "").split()) or f"Столбец {index}"
        count = used.get(name, 0) + 1
        used[name] = count
        headers.append(name if count == 1 else f"{name} {count}")
    return headers


def _numeric_cell(text: str) -> bool:
    value = (text or "").strip()
    if not value:
        return False
    if value.casefold() in _BOOLISH:
        return True
    try:
        return math.isfinite(float(value.replace(" ", "").replace(",", ".")))
    except ValueError:
        return False


def _column_kind(cells: list[str]) -> str:
    filled = [cell for cell in cells if not _skip_value(cell)]
    if not filled:
        return "skip"
    numeric = sum(1 for cell in filled if _numeric_cell(cell))
    if numeric >= 0.9 * len(filled):
        return "number"
    unique = {cell.casefold() for cell in filled}
    if len(unique) > 200:
        return "id"
    if len(unique) >= max(40, int(0.85 * len(filled))):
        return "id"
    return "category"


def _large_sheet(frame: pd.DataFrame, merges: list[tuple[int, int, int, int]], sheet_name: str) -> ExcelTable | None:
    if frame is None or frame.empty:
        return None
    work = frame.copy()
    _fill_merges(work, merges)
    grid = _grid(work)
    if len(grid) < 2:
        return None
    headers = _unique_headers(grid[0])
    body = grid[1:]
    width = len(headers)
    aligned = []
    for row in body:
        piece = list(row[:width]) + [""] * max(0, width - len(row))
        if any(not _skip_value(cell) for cell in piece):
            aligned.append(piece)
    if not aligned:
        return None
    kinds = []
    values_by_column: dict[str, list[str]] = {}
    for index, header in enumerate(headers):
        if _skip_header(header):
            kinds.append("skip")
            continue
        column = [row[index] for row in aligned]
        kind = _column_kind(column)
        kinds.append(kind)
        if kind != "category":
            continue
        seen = set()
        unique = []
        for cell in column:
            if _skip_value(cell):
                continue
            key = cell.casefold()
            if key in seen:
                continue
            seen.add(key)
            unique.append(cell)
        values_by_column[header] = unique
    return ExcelTable(
        sheet_name=sheet_name,
        headers=headers,
        kinds=kinds,
        rows=aligned,
        values_by_column=values_by_column,
    )


def _sheet_text(frame: pd.DataFrame, merges: list[tuple[int, int, int, int]]) -> str:
    if frame is None or frame.empty:
        return ""
    frame = frame.copy()
    _fill_merges(frame, merges)
    grid = _grid(frame)
    blocks = _blocks(grid)
    rendered = [text for text in (_render_block(block) for block in blocks) if text]
    if len(rendered) <= 1:
        return rendered[0] if rendered else ""
    return "\n\n".join(f"Таблица {index}\n{text}" for index, text in enumerate(rendered, 1))


def _fill_merges(frame: pd.DataFrame, merges: list[tuple[int, int, int, int]]) -> None:
    height, width = frame.shape
    for top, left, bottom, right in merges:
        if top >= height or left >= width or _missing(frame.iat[top, left]):
            continue
        value = frame.iat[top, left]
        for row in range(top, min(bottom, height - 1) + 1):
            for col in range(left, min(right, width - 1) + 1):
                if _missing(frame.iat[row, col]):
                    frame.iat[row, col] = value


def _grid(frame: pd.DataFrame) -> list[list[str]]:
    grid = [[_cell(value) for value in row] for row in frame.itertuples(index=False, name=None)]
    return _trim_border(grid)


def _trim_border(grid: list[list[str]]) -> list[list[str]]:
    """Срезает пустую кромку листа. Пустой столбец внутри не трогает: это граница таблиц."""
    if not grid:
        return []
    used_rows = [index for index, row in enumerate(grid) if any(row)]
    if not used_rows:
        return []
    grid = grid[used_rows[0]:used_rows[-1] + 1]
    width = max(len(row) for row in grid)
    used_cols = [
        col for col in range(width)
        if any(col < len(row) and row[col] for row in grid)
    ]
    if not used_cols:
        return []
    left, right = used_cols[0], used_cols[-1]
    width = right - left + 1
    trimmed = []
    for row in grid:
        piece = row[left:right + 1]
        if len(piece) < width:
            piece = piece + [""] * (width - len(piece))
        trimmed.append(piece)
    return trimmed


def _missing(value) -> bool:
    if value is None:
        return True
    try:
        return bool(pd.isna(value))
    except TypeError:
        return False


def _cell(value) -> str:
    if _missing(value):
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, datetime):
        if value.hour or value.minute or value.second:
            return value.strftime("%d.%m.%Y %H:%M")
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and math.isfinite(value):
        return _number(value)
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    text = str(value).replace("\r", " ").replace("\n", " ").strip()
    return " ".join(text.split())


def _number(value: float) -> str:
    rounded = round(value, 2 if abs(value) >= 1000 else 4)
    if abs(rounded - round(rounded)) < 1e-9:
        return str(int(round(rounded)))
    return f"{rounded:.4f}".rstrip("0").rstrip(".")


def _drop_empty(grid: list[list[str]]) -> list[list[str]]:
    if not grid:
        return []
    kept_rows = [row for row in grid if any(cell for cell in row)]
    if not kept_rows:
        return []
    width = max(len(row) for row in kept_rows)
    kept_cols = [col for col in range(width) if any(col < len(row) and row[col] for row in kept_rows)]
    if not kept_cols:
        return []
    return [[row[col] if col < len(row) else "" for col in kept_cols] for row in kept_rows]


def _blocks(grid: list[list[str]]) -> list[list[list[str]]]:
    """Прямоугольники, разделённые пустой строкой или пустым столбцом."""
    if not grid:
        return []
    found: list[list[list[str]]] = []
    for top, bottom in _spans(grid, by_row=True):
        band = grid[top:bottom]
        for left, right in _spans(band, by_row=False):
            piece = _drop_empty([row[left:right] for row in band])
            if not piece:
                continue
            nested = _blocks(piece) if piece != grid else []
            if nested and nested != [piece]:
                found.extend(nested)
            else:
                found.append(piece)
    return found


def _spans(grid: list[list[str]], by_row: bool) -> list[tuple[int, int]]:
    if not grid:
        return []
    limit = len(grid) if by_row else max(len(row) for row in grid)

    def empty(index: int) -> bool:
        if by_row:
            return not any(grid[index])
        return all(index >= len(row) or not row[index] for row in grid)

    spans = []
    start = None
    for index in range(limit):
        if empty(index):
            if start is not None:
                spans.append((start, index))
                start = None
        elif start is None:
            start = index
    if start is not None:
        spans.append((start, limit))
    return spans


def _render_block(block: list[list[str]]) -> str:
    if not block:
        return ""
    header_rows = _header_row_count(block)
    if header_rows == 0:
        headers = [""] * len(block[0])
        body = block
    else:
        headers = _compose_headers(block[:header_rows])
        body = [row for row in block[header_rows:] if any(row)]
        body = [row for row in body if row != headers]
    if not body:
        cells = [cell for cell in headers if cell]
        return " · ".join(cells)
    if header_rows == 0 and len(body) == 1:
        cells = [cell for cell in body[0] if cell]
        return " · ".join(cells)
    return _markdown(headers, body)


def _header_row_count(block: list[list[str]]) -> int:
    """1–3 верхние текстовые строки. Групповой заголовок реже, чем строка под ним."""
    if not block:
        return 0
    if _row_kind(block[0]) != "text":
        return 0
    count = 1
    while count < min(3, len(block)):
        current = block[count - 1]
        nxt = block[count]
        kind = _row_kind(nxt)
        if kind == "data":
            break
        if kind != "text":
            break
        filled = sum(1 for cell in current if cell)
        filled_next = sum(1 for cell in nxt if cell)
        if filled < filled_next:
            count += 1
            continue
        break
    return count


def _row_kind(row: list[str]) -> str:
    values = [cell for cell in row if cell]
    if not values:
        return "empty"
    numeric = sum(1 for cell in values if _numeric_text(cell))
    if numeric >= 1 and numeric >= len(values) - numeric:
        return "data"
    return "text"


def _numeric_text(value: str) -> bool:
    text = value.replace(" ", "").replace(",", ".")
    if text.endswith("%"):
        text = text[:-1]
    if not text:
        return False
    try:
        return math.isfinite(float(text))
    except ValueError:
        return False


def _compose_headers(rows: list[list[str]]) -> list[str]:
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    filled = [_spread_group_header(row) if _sparse(row) else row for row in normalized]
    headers = []
    for col in range(width):
        parts = []
        for row in filled:
            value = row[col]
            if value and (not parts or parts[-1] != value):
                parts.append(value)
        headers.append(" / ".join(parts))
    return headers


def _sparse(row: list[str]) -> bool:
    filled = sum(1 for cell in row if cell)
    return 0 < filled < len(row) * 0.6


def _spread_group_header(row: list[str]) -> list[str]:
    """Пустые ячейки объединённого заголовка получают название группы слева."""
    carried = ""
    spread = []
    for cell in row:
        if cell:
            carried = cell
            spread.append(cell)
        else:
            spread.append(carried)
    return spread


def _markdown(headers: list[str], body: list[list[str]]) -> str:
    width = max([len(headers), *[len(row) for row in body]])
    head = headers + [""] * (width - len(headers))

    def line(cells: list[str]) -> str:
        padded = cells + [""] * (width - len(cells))
        return "| " + " | ".join(_md(cell) for cell in padded) + " |"

    rows = [
        line(head),
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    rows.extend(line(row) for row in body)
    return "\n".join(rows)


def _md(value: str) -> str:
    return (value or "").replace("|", "\\|")
