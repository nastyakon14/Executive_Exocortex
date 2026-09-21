"""Конвертация Word → текст на macOS: python-docx. COM/Word здесь нет."""
from __future__ import annotations

from pathlib import Path


def convert_to_pdf(input_file, output_file=None):
    raise RuntimeError(
        "Конвертация Word → PDF через COM есть только на Windows. "
        "На macOS используйте .docx — его читает python-docx."
    )


def _extract_docx(input_path: str) -> str:
    import docx

    document = docx.Document(input_path)
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_word(input_path):
    """`.docx` читается напрямую. Старый `.doc` на Mac нужно сохранить как `.docx`."""
    ext = Path(input_path).suffix.lower()
    if ext == ".doc":
        raise RuntimeError(
            "Старый формат .doc на macOS не читается. Сохраните файл как .docx."
        )
    if ext != ".docx":
        raise RuntimeError(f"Ожидался .doc или .docx, получен {ext or 'файл без расширения'}")
    try:
        text = _extract_docx(input_path)
    except Exception as e:
        raise RuntimeError(f"Не удалось прочитать .docx: {e}") from e
    if not (text or "").strip():
        raise RuntimeError("В документе нет текста")
    return text
